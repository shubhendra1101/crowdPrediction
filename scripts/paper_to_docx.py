"""Convert docs/paper/main.tex into a Word file with plain, paste-ready text (no LaTeX code).

    python scripts/paper_to_docx.py --tex docs/paper/main.tex --out docs/paper/CrowdSafe_paper.docx

Handles exactly the LaTeX used in the paper: title/authors/abstract/keywords, sections and subsections
(IEEE numbering), enumerate lists, tables (caption "TABLE I: Name" above), figures (caption "Fig. 1. Name"
below), \\cite -> [n], \\ref -> table/figure/section numbers, inline math -> Unicode with Word sub/superscripts,
and the inline reference list.
"""
from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

FONT = "Times New Roman"
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
MACROS = {r"\max": "max", r"\sum": "Σ", r"\odot": "⊙", r"\rho": "ρ", r"\tau": "τ", r"\cdot": "·", r"\geq": "≥",
          r"\times": "×", r"\approx": "≈", r"\pm": "±", r"\dagger": "†", r"\,": " ", r"\;": " "}
# figure files in the paper -> raster versions for Word
FIG_SOURCES = {"pipeline.pdf": "results/figures/paper_pipeline.png", "dashboard.jpg": "docs/paper/figures/dashboard.jpg",
               "alert_timeline.pdf": "results/figures/paper_alert_timeline.png"}


def braced(s: str, i: int) -> tuple[str, int]:
    """Content of the {...} group starting at s[i] == '{' and the index after it."""
    assert s[i] == "{", s[i:i + 30]
    depth, j = 0, i
    while True:
        if s[j] == "{" and (j == 0 or s[j - 1] != "\\"):
            depth += 1
        elif s[j] == "}" and s[j - 1] != "\\":
            depth -= 1
            if depth == 0:
                return s[i + 1:j], j + 1
        j += 1


def math_runs(expr: str) -> list[tuple[str, str]]:
    """Inline math -> runs of (text, 'n'|'sub'|'sup')."""
    out: list[tuple[str, str]] = []
    i = 0
    while i < len(expr):
        c = expr[i]
        if c == "\\":
            m = re.match(r"\\(mathrm|mathcal|text)\{", expr[i:])
            if m:
                inner, j = braced(expr, i + len(m.group(0)) - 1)
                out.append(("𝒩" if (m.group(1) == "mathcal" and inner == "N") else inner, "n"))
                i = j
                continue
            m = re.match(r"\\[a-zA-Z]+|\\[,;]", expr[i:])
            tok = m.group(0)
            out.append((MACROS.get(tok, tok.lstrip("\\")), "n"))
            i += len(tok)
        elif c in "_^":
            kind = "sub" if c == "_" else "sup"
            if expr[i + 1] == "{":
                inner, j = braced(expr, i + 1)
            elif expr[i + 1] == "\\":
                m = re.match(r"\\[a-zA-Z]+", expr[i + 1:])
                inner, j = m.group(0), i + 1 + len(m.group(0))
            else:
                inner, j = expr[i + 1], i + 2
            out += [(t, kind) for t, _ in math_runs(inner)]
            i = j
        elif c in "{}":
            i += 1
        else:
            out.append((c, "n"))
            i += 1
    return out


class Converter:
    def __init__(self, tex: str):
        self.tex = tex
        bib = tex[tex.index(r"\begin{thebibliography}"):tex.index(r"\end{thebibliography}")]
        self.bib = re.findall(r"\\bibitem\{([^}]+)\}\s*(.*)", bib)
        self.cite_no = {k: n + 1 for n, (k, _) in enumerate(self.bib)}
        self.labels = self._number_labels()

    def _number_labels(self) -> dict[str, str]:
        body = self.tex[self.tex.index(r"\maketitle"):]
        labels, sec, sub, tab, fig = {}, 0, 0, 0, 0
        for m in re.finditer(r"\\(section|subsection)\{|\\label\{([^}]+)\}|\\CStabcap\{[^}]*\}\{([^}]+)\}|\\CSfigcap\{[^}]*\}\{([^}]+)\}", body):
            if m.group(1) == "section":
                sec, sub, last = sec + 1, 0, ROMAN[sec]
            elif m.group(1) == "subsection":
                sub += 1
                last = f"{ROMAN[sec - 1]}-{chr(64 + sub)}"
            elif m.group(2):
                labels[m.group(2)] = last
            elif m.group(3):
                tab += 1
                labels[m.group(3)] = ROMAN[tab - 1]
            elif m.group(4):
                fig += 1
                labels[m.group(4)] = str(fig)
        return labels

    # ---- inline text -> runs (text, style) with style in {'n','i','b','sub','sup'} ----
    def runs(self, s: str) -> list[tuple[str, str]]:
        s = re.sub(r"\\texorpdfstring\{([^{}]*(?:\{[^{}]*\}[^{}]*)*)\}\{[^}]*\}", r"\1", s)
        s = re.sub(r"~?\\cite\{([^}]+)\}", lambda m: " " + ", ".join(f"[{self.cite_no[k.strip()]}]" for k in m.group(1).split(",")), s)
        s = re.sub(r"\\ref\{([^}]+)\}", lambda m: self.labels[m.group(1)], s)
        s = s.replace(r"\pmsq{}", "persons/m²").replace(r"\pmsq", "persons/m²")
        s = s.replace("``", "“").replace("''", "”").replace("---", "—").replace("--", "–")
        s = s.replace(r'\"u', "ü").replace(r'\"a', "ä").replace(r'\"o', "ö").replace(r"\%", "%").replace(r"\_", "_")
        s = s.replace(r"\&", "&").replace(r"\,", " ").replace("~", " ").replace(r"\ ", " ")
        s = re.sub(r"\\url\{([^}]+)\}", r"\1", s)
        out: list[tuple[str, str]] = []
        i = 0
        while i < len(s):
            if s[i] == "$":
                j = s.index("$", i + 1)
                out += math_runs(s[i + 1:j])
                i = j + 1
                continue
            m = re.match(r"\\(emph|textbf)\{", s[i:])
            if m:
                inner, j = braced(s, i + len(m.group(0)) - 1)
                style = "i" if m.group(1) == "emph" else "b"
                out += [(t, style if st == "n" else st) for t, st in self.runs(inner)]
                i = j
                continue
            m = re.match(r"\\(centering|scriptsize|footnotesize|small)\b", s[i:])
            if m:
                i += len(m.group(0))
                continue
            out.append((s[i], "n"))
            i += 1
        merged: list[tuple[str, str]] = []
        for t, st in out:
            if merged and merged[-1][1] == st:
                merged[-1] = (merged[-1][0] + t, st)
            else:
                merged.append((t, st))
        return [(re.sub(r" {2,}", " ", t), st) for t, st in merged]

    def plain(self, s: str) -> str:
        return "".join(t for t, _ in self.runs(s)).strip()


def add_runs(par, runs, size=10, base_bold=False, base_italic=False):
    for t, st in runs:
        r = par.add_run(t)
        r.font.name, r.font.size = FONT, Pt(size)
        r.bold = base_bold or st == "b"
        r.italic = base_italic or st == "i"
        r.font.subscript = st == "sub"
        r.font.superscript = st == "sup"
    return par


def para(doc, runs, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, bold=False, italic=False, space_after=4, indent=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.17)
    return add_runs(p, runs, size, bold, italic)


def cell_border(table):
    tbl = table._tbl
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom", "insideH"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl.tblPr.append(borders)


def add_table(doc, conv: Converter, block: str):
    cap = re.search(r"\\CStabcap\{([^}]*)\}\{([^}]+)\}", block)
    para(doc, [(f"TABLE {conv.labels[cap.group(2)]}: {conv.plain(cap.group(1))}", "n")], 9,
         WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    start = block.index(r"\begin{tabular}") + len(r"\begin{tabular}")
    _, start = braced(block, start)                  # skip the column specification, whatever it contains
    tab = block[start:block.index(r"\end{tabular}")]
    note = None
    if r"\multicolumn" in tab:                       # table note: \multicolumn{n}{spec}{text}
        k = tab.index(r"\multicolumn") + len(r"\multicolumn")
        _, k = braced(tab, k)
        _, k = braced(tab, k)
        note, _ = braced(tab, k)
        tab = tab[:tab.index(r"\multicolumn")]
    head_part, _, body_part = tab.partition(r"\midrule")
    head_part = head_part.replace(r"\toprule", "")
    body_part = body_part.replace(r"\bottomrule", "")

    def rows(part):
        out = []
        for row in re.split(r"\\\\", part):
            row = row.replace(r"\midrule", "").strip()
            if row:
                out.append([c.strip() for c in re.split(r"(?<!\\)&", row)])
        return out

    head, body = rows(head_part), rows(body_part)
    if len(head) > 1:                             # two header lines -> one
        head = [[(a + " " + b).strip() for a, b in zip(head[0], head[1])]]
    ncol = len(head[0])
    t = doc.add_table(rows=0, cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_border(t)
    for r_i, row in enumerate(head + body):
        cells = t.add_row().cells
        for c_i in range(ncol):
            txt = row[c_i] if c_i < len(row) else ""
            p = cells[c_i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            add_runs(p, conv.runs(txt), 8, base_bold=(r_i < len(head)))
    if note:
        para(doc, conv.runs(note), 8, WD_ALIGN_PARAGRAPH.LEFT, italic=True, space_after=8)
    else:
        para(doc, [("", "n")], 4, space_after=4)


def add_figure(doc, conv: Converter, block: str, root: Path):
    src = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]+)\}", block).group(1)
    wide = r"\begin{figure*}" in block
    doc.add_picture(str(root / FIG_SOURCES[src]), width=Inches(6.5 if wide else 3.4))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = re.search(r"\\CSfigcap\{(.*)\}\{([^}]+)\}", block)
    para(doc, [(f"Fig. {conv.labels[cap.group(2)]}. ", "n")] + conv.runs(cap.group(1)), 9, WD_ALIGN_PARAGRAPH.CENTER,
         space_after=10)


def build(tex_path: Path, out: Path, root: Path) -> None:
    tex = tex_path.read_text(encoding="utf-8")
    tex = re.sub(r"(?m)(?<!\\)%.*$", "", tex)
    conv = Converter(tex)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.75)
    sec.top_margin = sec.bottom_margin = Inches(0.75)
    st = doc.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(10)

    title, _ = braced(tex, tex.index(r"\title{") + 6)
    para(doc, [(conv.plain(title.replace("\\\\", " ")), "n")], 22, WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    auth, _ = braced(tex, tex.index(r"\author{") + 7)
    for block in auth.split(r"\and"):
        name = re.search(r"\\IEEEauthorblockN\{([^}]*)\}", block).group(1)
        aff = re.search(r"\\IEEEauthorblockA\{(.*)\}", block, re.S).group(1)
        lines = [x.strip() for x in aff.split("\\\\") if x.strip()]
        para(doc, [(name, "n")], 11, WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
        for ln in lines:
            para(doc, [(ln, "n")], 9, WD_ALIGN_PARAGRAPH.CENTER, italic=not ln.endswith(".edu") and "India" not in ln,
                 space_after=0)
        para(doc, [("", "n")], 4, space_after=2)

    abstract = tex[tex.index(r"\begin{abstract}") + 16:tex.index(r"\end{abstract}")].strip()
    para(doc, [("Abstract—", "b")] + [(t, "b" if s == "n" else s) for t, s in conv.runs(abstract)], 9, space_after=6)
    kw = tex[tex.index(r"\begin{IEEEkeywords}") + 20:tex.index(r"\end{IEEEkeywords}")].strip()
    para(doc, [("Index Terms—", "b"), (conv.plain(kw), "b")], 9, space_after=10)

    body = tex[tex.index(r"\end{IEEEkeywords}") + 18:tex.index(r"\begin{thebibliography}")]
    pattern = re.compile(r"\\section\{|\\subsection\{|\\begin\{enumerate\}.*?\\end\{enumerate\}|"
                         r"\\begin\{table\*?\}.*?\\end\{table\*?\}|\\begin\{figure\*?\}.*?\\end\{figure\*?\}", re.S)
    sec_no = sub_no = 0
    pos = 0

    def flush(text):
        for chunk in re.split(r"\n\s*\n", text):
            chunk = re.sub(r"\\label\{[^}]*\}", "", chunk).strip()
            if chunk:
                para(doc, conv.runs(" ".join(chunk.split())), 10, indent=True)

    for m in pattern.finditer(body):
        flush(body[pos:m.start()])
        tok = m.group(0)
        if tok == r"\section{":
            name, end = braced(body, m.end() - 1)
            sec_no, sub_no = sec_no + 1, 0
            para(doc, [(f"{ROMAN[sec_no - 1]}. {conv.plain(name).upper()}", "n")], 10, WD_ALIGN_PARAGRAPH.CENTER,
                 space_after=4)
            pos = end
            continue
        if tok == r"\subsection{":
            name, end = braced(body, m.end() - 1)
            sub_no += 1
            para(doc, [(f"{chr(64 + sub_no)}. {conv.plain(name)}", "n")], 10, WD_ALIGN_PARAGRAPH.LEFT, italic=True,
                 space_after=3)
            pos = end
            continue
        if tok.startswith(r"\begin{enumerate}"):
            items = re.split(r"\\item\s+", tok[len(r"\begin{enumerate}"):-len(r"\end{enumerate}")])[1:]
            for n, it in enumerate(items, 1):
                para(doc, [(f"{n}) ", "n")] + conv.runs(" ".join(it.split())), 10, space_after=2)
        elif tok.startswith(r"\begin{table"):
            add_table(doc, conv, tok)
        elif tok.startswith(r"\begin{figure"):
            add_figure(doc, conv, tok, root)
        pos = m.end()
    flush(body[pos:])

    para(doc, [("REFERENCES", "n")], 10, WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    for n, (_, text) in enumerate(conv.bib, 1):
        para(doc, [(f"[{n}] ", "n")] + conv.runs(text.strip()), 8, WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--tex", type=Path, default=Path("docs/paper/main.tex"))
    ap.add_argument("--out", type=Path, default=Path("docs/paper/CrowdSafe_paper.docx"))
    a = ap.parse_args()
    build(a.tex, a.out, Path(__file__).resolve().parents[1])
    print("wrote", a.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
